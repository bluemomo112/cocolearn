#!/usr/bin/env python3
"""
读取Word和Excel文档的脚本
"""
import sys
from docx import Document
try:
    import openpyxl
    has_openpyxl = True
except ImportError:
    has_openpyxl = False
    import pandas as pd

def read_excel(file_path):
    """读取Excel文件"""
    print(f"\n=== 读取Excel文件: {file_path} ===\n")

    if has_openpyxl:
        wb = openpyxl.load_workbook(file_path)
        for sheet_name in wb.sheetnames:
            print(f"\n工作表: {sheet_name}")
            ws = wb[sheet_name]
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                if any(cell is not None for cell in row):
                    print(f"行{row_idx}: {row}")
    else:
        df = pd.read_excel(file_path, sheet_name=None)
        for sheet_name, data in df.items():
            print(f"\n工作表: {sheet_name}")
            print(data.to_string())

def read_word(file_path):
    """读取Word文档"""
    print(f"\n=== 读取Word文档: {file_path} ===\n")

    doc = Document(file_path)

    # 读取段落
    print("文档段落:")
    for i, para in enumerate(doc.paragraphs, 1):  # 读取所有段落
        if para.text.strip():
            print(f"\n段落{i}: {para.text}")

    # 读取表格
    if doc.tables:
        print(f"\n\n文档包含 {len(doc.tables)} 个表格")
        for table_idx, table in enumerate(doc.tables, 1):
            print(f"\n--- 表格 {table_idx} ---")
            for row_idx, row in enumerate(table.rows[:20], 1):  # 每个表格只显示前20行
                cells = [cell.text.strip() for cell in row.cells]
                print(f"行{row_idx}: {cells}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python read_docs.py <file_path>")
        sys.exit(1)

    file_path = sys.argv[1]

    if file_path.endswith('.xlsx') or file_path.endswith('.xls'):
        read_excel(file_path)
    elif file_path.endswith('.docx') or file_path.endswith('.doc'):
        read_word(file_path)
    else:
        print(f"不支持的文件格式: {file_path}")
